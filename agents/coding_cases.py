"""Reading and grouping the manual test cases the Coding Agent turns into code: the sheet, the cycle filter, chunking and split-file naming.

Split out of agents/testcase_writer.py, which re-exports every name below, so
`agents.testcase_writer.<name>` keeps working. Code is moved verbatim.
"""


from __future__ import annotations

import re
from pathlib import Path

import openpyxl

from agents.artifact_io import normalize_target_file, read_json
from agents.ts_check import _POM_PATH_MARKERS


def _load_planner_plans(planner_doc_path: Path) -> list:
    """Reads planner_doc.json into a list of plan dicts, or [] on anything
    unreadable/unexpected. Single loader shared by the three consumers below
    (re-inspection triggering, target_file grouping, per-chunk slicing) —
    they previously each opened and parsed the file with their own
    inline error handling.

    `target_file` is normalized HERE, at the reader, not only at the contract
    boundary. `qa_pipeline.agents.adapters.normalize_plan` does the same
    normalization, but it lands in automation_plan.json, and this file is
    what THIS agent reads — written raw by `automation_planner.run()`
    (including from its own `_cli`, with no orchestrator anywhere in the
    loop). So the two sides disagreed about what a file was called on every
    ordinary run: `_group_cases_by_target_file` keys on this exact string,
    which made ".\\tests\\a.spec.ts" and "tests/a.spec.ts" two different
    files — two generation calls, two versions of one spec, last write wins.
    See artifact_io.normalize_target_file.
    """
    planner_doc = read_json(planner_doc_path, default=[])
    plans = []
    for plan in planner_doc:
        if not isinstance(plan, dict):
            continue
        if plan.get("target_file"):
            # Copied, not mutated in place: read_json hands back a fresh
            # object each call, but a caller passing plans around shouldn't
            # find its input rewritten underneath it.
            plan = {**plan, "target_file": normalize_target_file(plan["target_file"])}
        plans.append(plan)
    return plans


def fetch_test_cycle_keys(cycle_id: str) -> list:
    """Fetches the test case keys in a Zephyr Essential/Scale test cycle.

    Requires ZEPHYR_API_TOKEN in your environment (see .env.example). Uses
    the api.zephyrscale.smartbear.com/v2 REST API (shared by Zephyr Essential
    and Zephyr Scale Cloud) — GET /testexecutions?testCycle=<key>, returning
    each execution's testCase.key. Falls back to "no filter" (empty list)
    with a clear warning if the token is missing or the cycle can't be
    reached, so the pipeline doesn't crash — it just runs against every case
    on the sheet instead of the ones actually in-cycle.

    NOTE: this endpoint shape (/testexecutions?testCycle=) is documented for
    Zephyr Scale Cloud's public API; verify it against your org's own Zephyr
    Essential API docs (usually linked from within the app itself) before
    trusting the filter in a real run — Zephyr Essential's docs weren't
    reachable to confirm this exact path at the time this was wired in.
    """
    import os

    import requests

    if not cycle_id:
        return []

    token = os.environ.get("ZEPHYR_API_TOKEN", "")
    if not token:
        print("⚠️  ZEPHYR_API_TOKEN missing — no test-cycle filter applied "
              "(all cases on the sheet will be used). Set it in .env to filter by cycle.")
        return []

    try:
        resp = requests.get(
            "https://api.zephyrscale.smartbear.com/v2/testexecutions",
            headers={"Authorization": f"Bearer {token}"},
            params={"testCycle": cycle_id, "maxResults": 1000},
            timeout=15,
        )
        resp.raise_for_status()
        executions = resp.json().get("values", [])
        return [e["testCase"]["key"] for e in executions if e.get("testCase", {}).get("key")]
    except (requests.RequestException, KeyError) as e:
        print(f"⚠️  Could not fetch test cycle '{cycle_id}' from Zephyr ({e}) — "
              "no filter applied, using all cases on the sheet.")
        return []


def read_test_case_sheet(sheet_path: Path) -> list:
    wb = openpyxl.load_workbook(sheet_path, data_only=True)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    rows = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        # A fully-blank row is not a test case. openpyxl counts a trailing row
        # that only ever held formatting (or was cleared by hand in Excel)
        # toward ws.max_row, and every consumer of this function then had to
        # cope with a case that is all None:
        #   - the review gate demanded a "Reviewed" mark on a row
        #     that can never carry one, making the gate unsatisfiable;
        #   - jira_uploader's own gate did the same, and had it passed would
        #     have filed a junk "(untitled test case)" issue for the blank row.
        # filter_automatable_cases and case_index.build happened to drop these
        # already (no "Automated" marker, no Test Case ID), which is why the
        # defect only surfaced at the two gates. Dropped once, here, rather
        # than re-guarded in each caller.
        if all(cell is None or (isinstance(cell, str) and not cell.strip()) for cell in row):
            continue
        # strict=True: a row with a different cell count than the header row
        # (a hand-edited sheet with a stray merged/deleted cell) would
        # otherwise zip() to whichever is shorter, silently misaligning every
        # column after the gap instead of raising — a malformed sheet should
        # fail loudly here, not produce test cases with shifted field values.
        try:
            rows.append(dict(zip(headers, row, strict=True)))
        except ValueError:
            raise ValueError(
                f"'{sheet_path}' row {row_idx} has {len(row)} cell(s) but the header "
                f"row has {len(headers)} — sheet is malformed (a hand-edited merged/"
                f"deleted cell is the usual cause). Fix the sheet before re-running."
            ) from None
    return rows


def filter_automatable_cases(cases: list, include_manual: bool = False) -> list:
    """Only rows marked 'Execution Type' == 'Automated' get planned/coded by
    default — the Test Case Writer's own skill (skills/testcase_writer.skill.md)
    defaults most UI-heavy cases to 'Manual', and generating Playwright code
    for a case the sheet itself calls manual silently overrides that human
    judgment call. Pass include_manual=True (`qa-pipeline run
    --include-manual`) to automate every row regardless of this column,
    for teams that want full automation coverage anyway."""
    if include_manual:
        return cases
    return [c for c in cases if (c.get("Execution Type") or "").strip().lower() == "automated"]


#: Fallback bucket key for cases the Planner doc has no target_file for
#: (missing planner doc, or a case id the Planner didn't cover).
_UNASSIGNED_BUCKET = "(unassigned)"


#: Hard cap on cases per LLM call regardless of target_file grouping.
#: Observed in practice: even 4 cases' worth of REAL Playwright code (spec +
#: POM files, imports, assertions) truncated a 16000-token response once the
#: model stopped writing placeholder content and started writing real code.
#: This is the floor for a repo with NO reusable POM summary — see
#: _max_cases_per_call() below for why a bigger existing_pom_summary raises it.
_MAX_CASES_PER_CALL = 2


def _max_cases_per_call(existing_pom_summary: str) -> int:
    """Raises the per-call case cap when there's a substantial POM summary
    to reuse from. The 16000-token OUTPUT budget is what truncates, not the
    input — but the amount of NEW code the model has to write per case is
    what actually consumes that output budget, and a bigger POM summary
    means more of each case's flow is "call ExistingPage.method(...)"
    instead of freshly-invented locators/assertions. Thresholds are a coarse
    heuristic (not measured against real truncation data at each tier the
    way _MAX_CASES_PER_CALL=2 was), so treat 3/4 as a starting point to
    tune against your own repo's POM_SUMMARY.md size, not a hard guarantee."""
    if not existing_pom_summary:
        return _MAX_CASES_PER_CALL
    length = len(existing_pom_summary)
    if length >= 5000:
        return 4
    if length >= 2000:
        return 3
    return _MAX_CASES_PER_CALL


#: A real Playwright test file is never this short; anything under this is
#: almost certainly placeholder/elided content, not a genuine (if minimal) file.
_MIN_REAL_CONTENT_LENGTH = 30


#: Exact strings observed in practice as literal placeholder "content" values.
_PLACEHOLDER_CONTENT_MARKERS = {"...", "…", "TODO", "full POM content"}


#: Matches Playwright's own default testMatch suffix so split-part filenames
#: still get collected as tests.
_TEST_SUFFIX_RE = re.compile(r"\.(spec|test)\.([cm]?[jt]sx?)$", re.I)


#: Matches a Playwright test/it title string: test('...', ...) or it("...", ...).
#: Deliberately permissive about quote style and whitespace — this is
#: checking a CONVENTION the skill asks for (skills/coding_agent.skill.md:
#: "Every test title MUST begin with its Test Case ID exactly as given"),
#: not parsing real JS/TS syntax.
_TEST_TITLE_RE = re.compile(r"""\b(?:test|it)\s*\(\s*['"`]([^'"`]*)""")


def _missing_title_prefixes(rel_path: str, content: str, expected_case_ids: list) -> list:
    """Returns the subset of expected_case_ids whose Test Case ID does NOT
    appear as the start of any test/it title in content — the convention
    every downstream agent's case-to-file resolution partially relies on
    (Self-Healer/Bug Reporter's title-prefix matching against the
    Executor's JSON report). Coding Agent's own 'test_case_ids' field on
    each file (the `generated_file` contract's test_case_ids) is the PRIMARY join,
    so a miss here is a real convention drift worth logging, not a reason
    to reject an otherwise-working file — hence a warning log, not
    rejection (see run()).

    Only test specs are checked. A page object carries the same
    `test_case_ids` as the spec it serves (that is how the two are joined),
    but it declares no test() at all, so every POM would otherwise report
    every one of its case ids as "missing" — a warning that fires on every
    run and can never be acted on. Same spec-only guard as
    _inline_locators_in_spec."""
    if not expected_case_ids or not _is_test_spec(rel_path):
        return []
    titles = _TEST_TITLE_RE.findall(content)
    return [cid for cid in expected_case_ids
            if cid and not any(title.startswith(cid) for title in titles)]


def _is_test_spec(rel_path: str) -> bool:
    normalized = str(rel_path or "").replace("\\", "/").lower()
    if any(marker in normalized for marker in _POM_PATH_MARKERS):
        return False
    return bool(_TEST_SUFFIX_RE.search(normalized))


def _part_label(chunk_cases: list, part: int) -> str:
    """A split file's name marker: the case ids it covers ("tc014-019"), not a
    running counter. Chunks finish in thread-pool order, so "_part3" meant a
    different file on every run and said nothing about what was inside it."""
    nums = []
    for case in chunk_cases or []:
        m = re.search(r"(\d+)\s*$", str(case.get("Test Case ID") or ""))
        if m:
            nums.append(m.group(1))
    if not nums:
        return f"part{part}"
    return f"tc{nums[0]}" + (f"-{nums[-1]}" if len(nums) > 1 and nums[-1] != nums[0] else "")


def _insert_part_suffix(path_str: str, part, label: str = "") -> str:
    """Renames a split-off file so Playwright still recognizes it as a test.

    Playwright's default testMatch pattern requires a file to END in
    ".spec.ts" (or .test.ts/.js/etc) — appending "_part{n}" via
    `Path.with_stem` (which only strips the LAST extension) produces
    "name.spec_part1.ts", which does NOT match and gets silently excluded
    from the run ("No tests found"), confirmed in practice. Insert the part
    marker BEFORE ".spec.ts" instead: "name_part1.spec.ts".
    """
    m = _TEST_SUFFIX_RE.search(path_str)
    if m:
        return path_str[:m.start()] + f"_{label or f'part{part}'}" + path_str[m.start():]
    # Not a recognized test-file suffix (e.g. a POM helper) — plain filename
    # split is fine, nothing depends on it matching testMatch.
    p = Path(path_str)
    return str(p.with_stem(f"{p.stem}_{label or f'part{part}'}"))


def _group_cases_by_target_file(cases: list, planner_doc_path: Path) -> dict:
    """Groups sheet rows by the Planner's target_file for their Test Case ID.

    One LLM call per target_file, not one call for the whole sheet: a full
    Playwright file's worth of code per case is verbose, and a single call
    covering ~18 cases reliably exceeded even a 16000-token cap. Grouping by
    target_file (rather than an arbitrary fixed-size batch) means every
    batch's output is naturally one coherent file — no risk of two batches
    both trying to write different content to the same path.
    """
    # An unreadable/absent planner doc yields {} here, which falls every case
    # through to the unassigned bucket — same behavior as before.
    target_by_id = {
        p["test_case_id"]: p["target_file"]
        for p in _load_planner_plans(planner_doc_path)
        if p.get("test_case_id") and p.get("target_file")
    }

    groups = {}
    for case in cases:
        target = target_by_id.get(case.get("Test Case ID")) or _UNASSIGNED_BUCKET
        groups.setdefault(target, []).append(case)
    return groups
