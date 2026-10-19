"""The test-case spreadsheet: its columns, writing a sheet, and reading an existing one.

Split out of agents/testcase_writer.py, which re-exports every name below, so
`agents.testcase_writer.<name>` keeps working. Code is moved verbatim.
"""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException

# Your manual test case format, in order — exactly as given, plus the
# trailing "Result" column found in the team's existing sheet (last run
# outcome — Pass/Fail/Not Executed).
MANUAL_COLUMNS = [
    "Name", "Description", "Product", "Component/s", "Sub-Component",
    "Prerequisite", "Test Step", "Test Data", "Expected Result",
    "Execution Type", "Testing Type", "Result",
]


# Human review gate column — the orchestrator halts before the Automation
# Planner runs unless every row here reads "Reviewed". Left blank by default;
# a human sets it after eyeballing the generated sheet.
REVIEW_COLUMN = ["Status"]


# Two columns appended for pipeline traceability — NOT part of your original
# spec, added because downstream agents (Planner, Coding Agent) need a
# unique ID to reference each case, and a way to trace back to the JIRA
# story/epic it came from. Both are plain text so they don't disturb your
# existing format if you filter/sort on the columns above.
PIPELINE_COLUMNS = ["Test Case ID", "JIRA Ref (Epic/Story)"]


COLUMNS = MANUAL_COLUMNS + REVIEW_COLUMN + PIPELINE_COLUMNS


#: Excel's hard per-cell character limit; openpyxl raises above this.
_EXCEL_CELL_LIMIT = 32767


def _to_cell_text(value):
    """Coerces a model-returned field value into Excel-writable text.

    The skill file asks for numbered-list strings (e.g. "1. ...\\n\\n2. ..."),
    but not every model follows that exactly — smaller/local models in
    particular return a JSON array of strings, or a list of step objects like
    [{"step": "...", "expected": "..."}]. openpyxl can't write a list/dict at
    all, and a raw Python repr in a cell is unreadable for the human who has
    to sign the sheet off at the review gate — so flatten recursively into the
    team's numbered-list house style instead.
    """
    if value is None:
        return ""

    if isinstance(value, list):
        parts = []
        for i, item in enumerate(value, start=1):
            text = _to_cell_text(item)
            # Only add "N-" numbering if the model didn't already number it.
            if text and not re.match(r"^\s*\d+\s*[-.)]", text):
                text = f"{i}-{text}"
            parts.append(text)
        rendered = "\n\n".join(p for p in parts if p)
    elif isinstance(value, dict):
        # Common step shapes: {"step"/"action": ..., "expected"/"result": ...}
        primary = next((value[k] for k in ("step", "action", "text", "description")
                        if value.get(k)), None)
        if primary:
            extra = next((value[k] for k in ("expected", "result", "expected_result")
                          if value.get(k)), None)
            rendered = f"{primary}" + (f" -> {extra}" if extra else "")
        else:
            rendered = "; ".join(f"{k}: {v}" for k, v in value.items())
    else:
        rendered = value

    if isinstance(rendered, str) and len(rendered) > _EXCEL_CELL_LIMIT:
        keep = _EXCEL_CELL_LIMIT - 20
        rendered = rendered[:keep] + "\n…[truncated]"
    return rendered


def _write_excel(cases: list, out_path: Path, epic_key: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Cases"

    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
    body_font = Font(name="Arial")
    wrap = Alignment(wrap_text=True, vertical="top")

    for col_idx, header in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Maps MANUAL_COLUMNS, in order, to the JSON keys the model returns
    # (see skills/testcase_writer.skill.md for the exact key names).
    key_map = [
        "name", "description", "product", "component", "sub_component",
        "prerequisite", "test_step", "test_data", "expected_result",
        "execution_type", "testing_type", "result",
    ]

    for row_idx, case in enumerate(cases, start=2):
        for col_idx, key in enumerate(key_map, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=_to_cell_text(case.get(key, "")))
            cell.font = body_font
            cell.alignment = wrap
        # Review gate column — blank until a human marks it "Reviewed",
        # EXCEPT when the case was imported from a sheet that already carried
        # that mark (see import_existing_test_cases). Re-asking for approval
        # a human has already given is friction, not safety.
        status_col = len(MANUAL_COLUMNS) + 1
        ws.cell(row=row_idx, column=status_col,
                value="Reviewed" if case.get("review_status") == "Reviewed" else ""
                ).font = body_font

        # Pipeline columns — auto-generated, not part of the model's output
        id_col = len(MANUAL_COLUMNS) + 2
        ref_col = len(MANUAL_COLUMNS) + 3
        ws.cell(row=row_idx, column=id_col,
                value=_to_cell_text(case.get("id")) or f"TC-{row_idx - 1:03d}").font = body_font
        ws.cell(row=row_idx, column=ref_col,
                value=f"{epic_key} / {case.get('_story_key', '?')}").font = body_font

    widths = [30, 32, 14, 18, 16, 24, 40, 18, 40, 14, 18, 14, 12, 14, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)  # 1_testcases/ is new per run
    wb.save(out_path)


#: OLE2 compound-document signature — the first 8 bytes of a genuine legacy
#: .xls (BIFF) workbook. A real OOXML workbook is a ZIP and starts "PK".
_OLE2_SIGNATURE = bytes.fromhex("d0cf11e0a1b11ae1")


def _load_workbook_tolerating_xls(path: Path):
    """openpyxl.load_workbook, but judging the file by its CONTENT rather
    than its extension.

    openpyxl refuses a `.xls` path outright, purely on the suffix, before
    looking inside. That matters because "an .xlsx someone renamed or
    saved with the wrong extension" is common in a folder of
    hand-maintained team spreadsheets — and for those files the refusal is
    wrong: the bytes are a perfectly readable OOXML workbook.

    So: try the normal path first, and if openpyxl rejects it, re-try
    through an in-memory buffer, which skips the extension check and lets
    openpyxl decide from the actual bytes. A genuine BIFF workbook still
    fails there, and only then is it reported.

    Either way the failure surfaces as one actionable ValueError. Both of
    the underlying errors are unhelpful on their own and neither was
    handled before: InvalidFileException for a .xls suffix, and
    zipfile.BadZipFile ("File is not a zip file") for a .xlsx that isn't
    actually OOXML — a CSV or HTML export saved under the wrong name, which
    is exactly the kind of file that turns up in a folder import.
    """
    import io
    import zipfile

    unreadable = (InvalidFileException, zipfile.BadZipFile)

    try:
        return openpyxl.load_workbook(path, data_only=True)
    except unreadable as exc:
        data = path.read_bytes()
        if not data.startswith(_OLE2_SIGNATURE):
            # Not a BIFF file, so the rejection may have been about the
            # extension alone — let openpyxl look at the bytes instead.
            try:
                return openpyxl.load_workbook(io.BytesIO(data), data_only=True)
            except unreadable:
                pass  # genuinely unreadable; fall through to the clear error
        is_biff = data.startswith(_OLE2_SIGNATURE)
        detail = ("It is a genuine legacy .xls (BIFF) workbook, which is not supported "
                  "— openpyxl reads only .xlsx/.xlsm and this project does not depend "
                  "on xlrd." if is_biff else
                  "Its contents are not an .xlsx/.xlsm workbook, whatever the file is "
                  "named (a CSV or HTML export saved with a spreadsheet extension does "
                  "this).")
        raise ValueError(
            f"'{path}' cannot be read as a spreadsheet. {detail} Open it in Excel or "
            f"LibreOffice and re-save as .xlsx, then re-run. (openpyxl said: {exc})"
        ) from exc


def _read_existing_test_case_file(path: Path) -> list:
    """Reads an existing team-authored test case file in WHATEVER shape it's
    already in — no assumption its columns match this pipeline's canonical
    MANUAL_COLUMNS, since normalizing a non-matching format is the entire
    point of import mode. Supports .xlsx/.xlsm (first/active sheet) and
    .csv. A row is kept if at least one cell has content; fully blank rows
    (common at the end of a hand-maintained sheet) are skipped.

    A .xls SUFFIX is accepted, but what happens next depends on the file's
    actual CONTENT, not its name — see _load_workbook_tolerating_xls. A
    genuine legacy .xls (BIFF) cannot be read at all (openpyxl supports only
    the OOXML formats and `xlrd` is not a dependency here), and used to
    escape as a raw openpyxl InvalidFileException while this docstring, the
    unsupported-format error below, and the --import-testcases help all
    advertised .xls as supported.
    """
    suffix = path.suffix.lower()

    if suffix in (".xlsx", ".xlsm", ".xls"):
        wb = _load_workbook_tolerating_xls(path)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h).strip() if h is not None else f"column_{i}"
                   for i, h in enumerate(rows[0])]
        out = []
        for row in rows[1:]:
            if not any(c is not None and str(c).strip() for c in row):
                continue
            out.append({headers[i]: row[i] for i in range(min(len(headers), len(row)))})
        return out

    if suffix == ".csv":
        import csv
        with path.open(newline="", encoding="utf-8-sig") as f:
            return [row for row in csv.DictReader(f) if any((v or "").strip() for v in row.values())]

    raise ValueError(
        f"Unsupported existing-test-case file format '{suffix}' for '{path}' — "
        f"expected .xlsx, .xlsm, or .csv (a legacy .xls must be re-saved as .xlsx "
        f"first; see _read_existing_test_case_file)."
    )
