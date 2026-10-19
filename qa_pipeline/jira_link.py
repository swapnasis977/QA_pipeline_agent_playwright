"""`qa-pipeline link-jira`: tie every automated test to the Jira test it was uploaded as.

The uploader creates (or updates) a Jira Test per sheet row and records `case id -> Jira key` in
`jira_upload_result.json`. Nothing carried those keys back, so the sheet said `PROJ-68294-AUTO-TC-107`
and Jira said `PROJ-71234`, and nobody could go from one to the other. This writes the keys back:

  * a `Jira Key` column in the workbook(s), found by Test Case ID (or Automation ID);
  * `jira-map.json` at the workspace root, which the `linkJira` test fixture reads so every test in
    the report carries a `jira` annotation naming its Jira test.

It only READS the upload result; it never calls Jira.
"""

from __future__ import annotations

import json
import re
from pathlib import Path

COLUMN = "Jira Key"
_BACKUP = re.compile(r"\.backup|_BACKUP", re.I)


def keys_from_results(results_path) -> dict[str, str]:
    """{case id: Jira key} from `jira_upload_result.json`; rows that failed to upload are left out."""
    try:
        rows = json.loads(Path(results_path).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    return {str(r["test_case_id"]): str(r["jira_key"]) for r in rows
            if isinstance(r, dict) and r.get("test_case_id") and r.get("jira_key")}


def write_map(workspace, keys: dict[str, str]) -> Path:
    """`jira-map.json` at the workspace root, merged with what is already there."""
    path = Path(workspace) / "jira-map.json"
    try:
        known = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        known = {}
    known.update(keys)
    path.write_text(json.dumps(dict(sorted(known.items())), indent=2) + "\n", encoding="utf-8")
    return path


def write_sheet_column(sheet_path, keys: dict[str, str]) -> int:
    """Adds / fills the `Jira Key` column; returns how many rows got a key."""
    import openpyxl
    from openpyxl.utils import get_column_letter

    workbook = openpyxl.load_workbook(sheet_path)
    filled = 0
    for sheet in workbook:
        header = [c.value for c in sheet[1]]
        if "Test Case ID" not in header:
            continue
        ix = {h: i for i, h in enumerate(header) if h}
        if COLUMN not in ix:
            col = len(header) + 1
            sheet.cell(1, col, COLUMN)
            sheet.cell(1, col)._style = sheet.cell(1, col - 1)._style
            sheet.column_dimensions[get_column_letter(col)].width = 14
            ix[COLUMN] = col - 1
        for row in sheet.iter_rows(min_row=2):
            ids = [str(row[ix["Test Case ID"]].value or "")]
            if "Automation ID" in ix:
                ids.append(str(row[ix["Automation ID"]].value or ""))
            key = next((keys[i] for i in ids if i in keys), None)
            if key:
                cell = sheet.cell(row[0].row, ix[COLUMN] + 1, key)
                cell._style = sheet.cell(row[0].row, 1)._style
                filled += 1
    workbook.save(sheet_path)
    return filled


def link(profile, results_path=None, sheets=None) -> dict:
    """Writes the keys back everywhere; returns what it did."""
    if results_path is None:
        found = sorted(Path(profile.context_root).glob("**/jira_upload_result.json"),
                       key=lambda p: p.stat().st_mtime)
        results_path = found[-1] if found else None
    keys = keys_from_results(results_path) if results_path else {}
    if not keys:
        return {"results": str(results_path or ""), "keys": 0, "map": "", "sheets": {}}
    if sheets is None:
        sheets = [p for p in Path(profile.context_root).glob("*/1_testcases/*.xlsx")
                  if not _BACKUP.search(p.name)]
    return {"results": str(results_path), "keys": len(keys),
            "map": str(write_map(profile.workspace, keys)),
            "sheets": {str(p): write_sheet_column(p, keys) for p in sheets}}


def cmd_link_jira(args, profile) -> int:
    report = link(profile, args.results, [Path(s) for s in args.sheet] if args.sheet else None)
    if not report["keys"]:
        print("      No Jira keys found. Upload first (the uploader writes jira_upload_result.json), "
              "or pass --results <that file>.")
        return 1
    print(f"      {report['keys']} Jira key(s) from {report['results']}")
    print(f"      Wrote {report['map']}  (the linkJira fixture reads it: each test in the report "
          f"then carries a `jira` annotation)")
    for path, filled in report["sheets"].items():
        print(f"      {filled:4d} row(s) given a '{COLUMN}' in {path}")
    return 0
