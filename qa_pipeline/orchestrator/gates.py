"""The two human touchpoints. Neither of them is authoring.

  1. APPROVE the test case sheet. Read the generated cases, tick Reviewed.
  2. DEMONSTRATE an unreachable screen. Click through it once under
     `playwright codegen` when the crawler cannot get there.

Everything else the pipeline does itself. It is worth being honest about why
touchpoint 2 cannot be removed entirely: hardware-gated states, external
redirects and seeded records genuinely cannot be crawled to, and the only
alternative to a person clicking through once is a model guessing — which is
the thing the whole architecture exists to prevent. Clicking is not writing a
script, and the resulting recording becomes the highest-confidence source in
the ledger, permanently, so the same screen is never asked for twice.

The review gate is ported from `run_pipeline._enforce_review_gate` and keeps
its two hard-won details:

  - A fully blank row is skipped. openpyxl counts a trailing row that only
    ever carried formatting toward `ws.max_row`, and such a row can never be
    marked Reviewed — so without the skip the gate became permanently
    unsatisfiable on any sheet someone had opened in Excel.
  - A header-only sheet HALTS rather than passing. Zero rows means zero
    unreviewed rows, which the naive check waves straight through.

It raises instead of calling `sys.exit`, so a caller embedding the pipeline
(a service, a test, a CI wrapper) can handle it rather than having the process
killed underneath it.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl


class ReviewRequired(Exception):
    """The sheet needs human approval before the pipeline may continue.

    Explicitly NOT an error. It is the pipeline stopping exactly where it is
    designed to stop, and the message says so — a human reading a red
    traceback assumes something broke and starts debugging.
    """

    def __init__(self, message: str, sheet_path=None, rows=None):
        super().__init__(message)
        self.sheet_path = sheet_path
        self.rows = list(rows or [])


class DemonstrationRequired(Exception):
    """Rung 4 of the escalation ladder: a screen the Explorer cannot reach.

    Carries everything the person needs to act without reading any code —
    which element, on which screen, and the exact command to record it.
    """

    def __init__(self, items, recordings_dir=None):
        self.items = list(items or [])
        self.recordings_dir = recordings_dir
        super().__init__(self._render())

    def _render(self) -> str:
        lines = [
            "",
            "Some elements could not be reached by crawling, so the pipeline needs",
            "one demonstration from you. This is clicking, not coding — and it is",
            "asked once per screen, ever.",
            "",
        ]
        for item in self.items:
            lines.append(f"  - {item.get('logical_name', '?')} "
                         f"on {item.get('screen', '?')}")
            if item.get("reason"):
                lines.append(f"      why: {item['reason']}")
        lines += [
            "",
            "  What to do:",
            "    1. npx playwright codegen <your app url>",
            "    2. Click through to the screen(s) above and interact with the",
            "       element(s) once.",
            f"    3. Save the generated file into {self.recordings_dir or 'the profile recordings/ folder'}",
            "    4. Re-run the exact same command. The recording becomes the",
            "       highest-confidence source in the ledger and is never asked for again.",
            "",
        ]
        return "\n".join(lines)


def enforce_review(sheet_path, *, required: bool = True):
    """Halts unless every test case row's Status column reads exactly
    'Reviewed'. The Test Case Writer leaves it blank on purpose."""
    sheet_path = Path(sheet_path)
    if not required:
        print("      Review gate disabled in this profile — proceeding without approval.")
        return

    if not sheet_path.is_file():
        raise ReviewRequired(
            f"[HALT] No test case sheet at {sheet_path}.", sheet_path)

    workbook = openpyxl.load_workbook(sheet_path, data_only=True)
    worksheet = workbook.active
    headers = [c.value for c in worksheet[1]]

    if "Status" not in headers:
        raise ReviewRequired(
            f"[HALT] '{sheet_path.name}' has no 'Status' column — the review gate "
            f"cannot be verified.", sheet_path)
    status_col = headers.index("Status")

    if worksheet.max_row < 2:
        raise ReviewRequired(
            f"[HALT] '{sheet_path.name}' has a header row but no test cases. "
            f"Re-run with --force testcase_writer to rebuild it.", sheet_path)

    unreviewed, case_rows = [], 0
    for row_idx, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        if all(cell is None or (isinstance(cell, str) and not cell.strip()) for cell in row):
            continue
        case_rows += 1
        cell = row[status_col] if status_col < len(row) else None
        status = cell.strip() if isinstance(cell, str) else cell
        if status != "Reviewed":
            unreviewed.append(row_idx)

    if not case_rows:
        raise ReviewRequired(
            f"[HALT] '{sheet_path.name}' has a header row but no test cases. "
            f"Re-run with --force testcase_writer to rebuild it.", sheet_path)

    if unreviewed:
        raise ReviewRequired(
            "\n".join([
                "",
                "[HALT] Stopping here on purpose — this is not an error.",
                f"       {len(unreviewed)} of {case_rows} test case(s) in "
                f"'{sheet_path.name}' are not yet marked reviewed.",
                f"       Row(s): {unreviewed}",
                "",
                "       A human must read each case before any code is generated from it.",
                "       Open the sheet, and for each case that looks right set its",
                "       'Status' column to exactly 'Reviewed' (case-sensitive).",
                "       Then run the same command again — it picks up where it left off.",
                "",
            ]),
            sheet_path, unreviewed)

    print(f"      Review gate passed: {case_rows} case(s) approved.")


def request_demonstrations(unresolved: list, recordings_dir):
    """Raises `DemonstrationRequired` when anything reached rung 4.

    Called once, at the end of planning, rather than per element — asking a
    person for five recordings in one message respects their time; asking five
    times across a forty-minute run does not.
    """
    if not unresolved:
        return
    raise DemonstrationRequired(unresolved, recordings_dir)
