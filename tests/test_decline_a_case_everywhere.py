"""Declining a case must change EVERY place it lives, or it comes back half-declined.

The failure this prevents: TC-065 ("Hide Sidebar") was declined by hand in the run JSON, the plan, the
spec and three workbooks. Miss one and the sheet says "Automated" while the script is skipped, or
the script still runs red against a feature the build does not have.
"""

from __future__ import annotations

import json
from types import SimpleNamespace

import openpyxl

from qa_pipeline import decline
from utils import final_suite

SPEC = """import { test, expect } from '@fixtures/index';

test(
  'PROJ-1-AUTO-TC-065 Verify Hide Sidebar collapses the panel',
  { tag: ['@newfeature'] },
  async ({ page }) => {
    test.setTimeout(240_000);
    await page.goto('/');
  }
);

test('PROJ-1-AUTO-TC-066 Another case', async ({ page }) => {
  await page.goto('/');
});
"""


def _profile(tmp_path):
    root = tmp_path / "profiles" / "app"
    run = tmp_path / "runs" / "app" / "FEATURE" / "1_testcases"
    auto = tmp_path / "runs" / "app" / "FEATURE" / "2_automation"
    tests = tmp_path / "workspaces" / "app" / "tests" / "e2e"
    for folder in (root, run, auto, tests):
        folder.mkdir(parents=True)
    (root / "automation_scope.json").write_text(json.dumps({"rules": [], "default": "new"}))
    cases = [{"test_case_id": "PROJ-1-AUTO-TC-065", "declined": False, "decline_reason": ""},
             {"test_case_id": "PROJ-1-AUTO-TC-066", "declined": False, "decline_reason": ""}]
    (run / "test_cases.json").write_text(json.dumps(cases))
    (auto / "automation_plan.json").write_text(json.dumps(cases))
    (tests / "panel.spec.ts").write_text(SPEC)
    sheet = openpyxl.Workbook()
    ws = sheet.active
    ws.append(["Test Case ID", "Automation ID", "Scope Decision", "Automation Status", "Decision Note",
               "Status", "Last Run Result", "Result"])
    ws.append(["PROJ-1-TC-009", "PROJ-1-AUTO-TC-065", "Already automated", "Automated - script exists",
               "", "Reviewed", "Failed", "Failed"])
    ws.append(["PROJ-1-TC-010", "PROJ-1-AUTO-TC-066", "Already automated", "Automated - script exists",
               "", "Reviewed", "Passed", "Passed"])
    sheet.save(run / "Master.xlsx")
    return SimpleNamespace(root=root, context_root=tmp_path / "runs" / "app",
                           workspace=tmp_path / "workspaces" / "app"), run


def test_one_command_declines_the_json_the_plan_the_spec_and_the_workbook(tmp_path):
    profile, run = _profile(tmp_path)
    report = decline.decline(profile, "PROJ-1-AUTO-TC-065", "Feature not in build; only Dock exists.")
    assert len(report["json"]) == 2 and len(report["specs"]) == 1 and len(report["sheets"]) == 1

    for name in ("1_testcases/test_cases.json", "2_automation/automation_plan.json"):
        entries = json.loads((run.parent / name).read_text())
        by_id = {e["test_case_id"]: e for e in entries}
        assert by_id["PROJ-1-AUTO-TC-065"]["declined"] is True
        assert by_id["PROJ-1-AUTO-TC-066"]["declined"] is False        # a neighbour is untouched

    spec = (tmp_path / "workspaces" / "app" / "tests" / "e2e" / "panel.spec.ts").read_text()
    first_test, second_test = spec.split("test('PROJ-1-AUTO-TC-066")
    assert "test.skip(true, 'Declined: Feature not in build; only Dock exists.');" in first_test
    assert "test.skip" not in second_test
    assert "await page.goto('/');" in first_test                      # the steps are kept

    rows = list(openpyxl.load_workbook(run / "Master.xlsx").active.iter_rows(min_row=2, values_only=True))
    assert rows[0][2:4] == (decline.STATUS, decline.STATUS) and rows[0][6] == decline.RESULT
    assert rows[1][2] == "Already automated"


def test_declining_twice_does_not_stack_skips(tmp_path):
    profile, _ = _profile(tmp_path)
    decline.decline(profile, "PROJ-1-AUTO-TC-065", "reason")
    again = decline.decline(profile, "PROJ-1-AUTO-TC-065", "reason")
    spec = (tmp_path / "workspaces" / "app" / "tests" / "e2e" / "panel.spec.ts").read_text()
    assert spec.count("test.skip(true") == 1 and again["specs"] == [] and again["scope_file"] is False


def test_a_dry_run_changes_nothing(tmp_path):
    profile, run = _profile(tmp_path)
    before = (tmp_path / "workspaces" / "app" / "tests" / "e2e" / "panel.spec.ts").read_text()
    report = decline.decline(profile, "PROJ-1-AUTO-TC-065", "reason", dry_run=True)
    assert report["specs"] and report["sheets"]
    assert (tmp_path / "workspaces" / "app" / "tests" / "e2e" / "panel.spec.ts").read_text() == before
    assert json.loads((run / "test_cases.json").read_text())[0]["declined"] is False


def test_a_case_built_in_a_loop_is_reported_for_a_hand_edit_not_edited_blindly(tmp_path):
    profile, _ = _profile(tmp_path)
    spec = tmp_path / "workspaces" / "app" / "tests" / "e2e" / "loop.spec.ts"
    spec.write_text("for (const [id] of [['PROJ-1-AUTO-TC-094', 'Grayscale']]) { test(`${id} x`, async () => {}); }\n")
    report = decline.decline(profile, "PROJ-1-AUTO-TC-094", "reason")
    assert report["specs_by_hand"] == [str(spec)] and report["specs"] == []


def test_a_reason_is_required(tmp_path):
    profile, _ = _profile(tmp_path)
    try:
        decline.decline(profile, "PROJ-1-AUTO-TC-065", "  ")
    except ValueError as exc:
        assert "reason" in str(exc)
    else:
        raise AssertionError("an empty reason must be refused")


def test_a_rebuilt_sheet_keeps_the_decision(tmp_path):
    profile, _ = _profile(tmp_path)
    decline.decline(profile, "PROJ-1-AUTO-TC-065", "Feature not in build.")
    scope = final_suite.load_scope(profile.root / "automation_scope.json")
    rows = [{"Test Case ID": "PROJ-1-TC-009", "Automation ID": "PROJ-1-AUTO-TC-065", "Name": "Hide sidebar",
             "Automation Status": "Automated - script exists"}]
    final_suite.annotate(rows, scope)
    assert rows[0]["Scope Decision"] == "Declined - not applicable"
    assert rows[0]["Decision Note"] == "Feature not in build."
