"""After an upload, every automated test must be traceable to its Jira test.

The failure this prevents: the sheet said `PROJ-68294-AUTO-TC-107`, Jira said `PROJ-71234`, and the
uploader's result file held the pairing where nobody looked. Without the keys written back, "which
Jira test is this script?" had no answer, and the Jira upload could not be mapped to the Bitbucket suite.
"""

from __future__ import annotations

import json
from types import SimpleNamespace

import openpyxl

from qa_pipeline import jira_link


def _setup(tmp_path):
    runs = tmp_path / "runs" / "app"
    feature = runs / "FEATURE"
    (feature / "1_testcases").mkdir(parents=True)
    results = feature / "jira_upload_result.json"
    results.write_text(json.dumps([
        {"test_case_id": "PROJ-1-AUTO-TC-001", "jira_key": "PROJ-9001", "action": "created"},
        {"test_case_id": "PROJ-1-TC-002", "jira_key": "PROJ-9002", "action": "updated"},
        {"test_case_id": "PROJ-1-AUTO-TC-003", "jira_key": None, "action": "error"},
    ]))
    sheet = feature / "1_testcases" / "Final.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Test Case ID", "Automation ID", "Name"])
    ws.append(["PROJ-1-AUTO-TC-001", "", "first"])
    ws.append(["PROJ-1-TC-009", "PROJ-1-TC-002", "second, listed by manual id with an automation id"])
    ws.append(["PROJ-1-AUTO-TC-003", "", "failed to upload"])
    wb.save(sheet)
    workspace = tmp_path / "workspaces" / "app"
    workspace.mkdir(parents=True)
    return SimpleNamespace(context_root=runs, workspace=workspace), results, sheet


def test_keys_go_into_the_sheet_and_into_the_workspace_map(tmp_path):
    profile, results, sheet = _setup(tmp_path)
    report = jira_link.link(profile, results)
    assert report["keys"] == 2                                   # the failed upload has no key to write
    rows = list(openpyxl.load_workbook(sheet).active.iter_rows(min_row=1, values_only=True))
    assert rows[0][-1] == jira_link.COLUMN
    assert [r[-1] for r in rows[1:]] == ["PROJ-9001", "PROJ-9002", None]
    assert json.loads((profile.workspace / "jira-map.json").read_text()) == {
        "PROJ-1-AUTO-TC-001": "PROJ-9001", "PROJ-1-TC-002": "PROJ-9002"}


def test_linking_twice_does_not_duplicate_the_column_or_lose_earlier_keys(tmp_path):
    profile, results, sheet = _setup(tmp_path)
    jira_link.link(profile, results)
    results.write_text(json.dumps([{"test_case_id": "PROJ-1-AUTO-TC-003", "jira_key": "PROJ-9003"}]))
    jira_link.link(profile, results)
    header = [c.value for c in openpyxl.load_workbook(sheet).active[1]]
    assert header.count(jira_link.COLUMN) == 1
    assert len(json.loads((profile.workspace / "jira-map.json").read_text())) == 3


def test_a_missing_upload_result_is_reported_not_crashed(tmp_path):
    profile, _, _ = _setup(tmp_path)
    assert jira_link.link(profile, tmp_path / "nope.json")["keys"] == 0
