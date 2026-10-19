"""The master sheet must tell the truth about what is automated.

The failure this prevents: the cases lived in four places with three numbering schemes (story
sheets, a consolidated sheet, the run's sheet, the spec files), so "how many test cases do we
have, and how many are automated?" had no single answer - the same feature was quoted as 543,
89, 48 and 43 depending on which file you opened. A sheet that over-claims ("automated" for a
case with no script) is worse than none.
"""

import openpyxl

from utils import master_sheet as ms


def _case(name, execution="Automated", test_id="S-1-TC-001", ref=""):
    return {"Name": name, "Execution Type": execution, "Test Case ID": test_id,
            "JIRA Ref (Epic/Story)": ref, "Test Step": "1. do it", "Expected Result": "1. ok"}


def _rows(per_story=None, consolidated=(), run_sheet=(), scripts=None, results=None):
    return ms.build_rows(per_story or {}, {"S-1": "Zoom"}, list(consolidated), list(run_sheet),
                         scripts or {}, results or {})


def _by_name(rows):
    return {r["Name"]: r for r in rows}


def test_a_case_with_a_script_is_automated_and_carries_its_file_and_result():
    rows = _rows({"S-1": [_case("Zoom in")]},
                 consolidated=[_case("Zoom in", test_id="AUTO-1", ref="F / S-1")],
                 scripts={"AUTO-1": "tests/e2e/zoom.spec.ts"},
                 results={"AUTO-1": (ms.PASSED, "")})
    row = rows[0]
    assert row["Automation Status"] == ms.AUTOMATED
    assert row["Automated Test File"] == "tests/e2e/zoom.spec.ts"
    assert row["Last Run Result"] == ms.PASSED


def test_a_picked_case_with_no_script_is_not_called_automated():
    rows = _rows({"S-1": [_case("Zoom in")]},
                 consolidated=[_case("Zoom in", test_id="AUTO-1", ref="F / S-1")])
    assert rows[0]["Automation Status"] == ms.SELECTED
    assert rows[0]["Last Run Result"] == ""          # no script, so no result to show


def test_an_unpicked_automatable_case_is_flagged_as_such():
    assert _rows({"S-1": [_case("Pan")]})[0]["Automation Status"] == ms.AUTOMATABLE


def test_a_manual_case_stays_manual_even_if_the_consolidation_picked_it():
    rows = _rows({"S-1": [_case("Visual check", execution="Manual")]},
                 consolidated=[_case("Visual check", execution="Manual", test_id="AUTO-9",
                                     ref="F / S-1")])
    assert rows[0]["Automation Status"] == ms.MANUAL


def test_a_script_with_no_entry_in_the_last_report_is_not_run_not_passed():
    rows = _rows({"S-1": [_case("Zoom in")]},
                 consolidated=[_case("Zoom in", test_id="AUTO-1", ref="F / S-1")],
                 scripts={"AUTO-1": "tests/e2e/zoom.spec.ts"}, results={})
    assert rows[0]["Last Run Result"] == ms.NOT_RUN


def test_cases_that_only_automation_added_get_their_own_rows_without_double_counting():
    run_sheet = [_case("Zoom in", test_id="AUTO-1"),                       # already a story case
                 _case("Brand new", test_id="AUTO-90", ref="F / S-7")]     # exists only here
    rows = _rows({"S-1": [_case("Zoom in")]},
                 consolidated=[_case("Zoom in", test_id="AUTO-1", ref="F / S-1")],
                 run_sheet=run_sheet, scripts={"AUTO-90": "tests/e2e/new.spec.ts"})
    assert len(rows) == 2
    assert _by_name(rows)["Brand new"]["Test Case ID"] == "AUTO-90"
    assert _by_name(rows)["Brand new"]["Automation Status"] == ms.AUTOMATED


def test_the_same_name_in_two_stories_is_matched_to_the_right_story():
    per_story = {"S-1": [_case("Reset", test_id="S-1-TC-001")], "S-2": [_case("Reset", test_id="S-2-TC-001")]}
    rows = _rows(per_story, consolidated=[_case("Reset", test_id="AUTO-5", ref="F / S-2")])
    by_story = {r["Story"]: r for r in rows}
    assert by_story["S-2"]["Automation ID"] == "AUTO-5"
    assert by_story["S-1"]["Automation ID"] == ""


REPORT = {"suites": [{"specs": [
    {"title": "PROJ-1-AUTO-TC-001 passes", "tests": [{"results": [{"status": "passed"}]}]},
    {"title": "PROJ-1-AUTO-TC-002 fails", "tests": [{"results": [{"status": "failed", "error": {
        "message": "\x1b[31mTimeoutError: locator.waitFor: Timeout 30000ms exceeded.\x1b[39m\nmore"}}]}]},
    {"title": "PROJ-1-AUTO-TC-003 cut off", "tests": [{"status": "skipped", "results": []}]},
    {"title": "PROJ-1-AUTO-TC-004 half", "tests": [{"results": [{"status": "passed"}]},
                                                {"results": [{"status": "failed"}]}]},
]}]}


def test_last_results_read_a_playwright_report_and_take_the_worst_outcome():
    got = ms.last_results(REPORT, {"PROJ-1-AUTO-TC-002": "stalled"})
    assert got["PROJ-1-AUTO-TC-001"][0] == ms.PASSED
    assert got["PROJ-1-AUTO-TC-003"][0] == ms.NOT_RUN
    assert got["PROJ-1-AUTO-TC-004"][0] == ms.FAILED          # one failing test fails the case
    outcome, note = got["PROJ-1-AUTO-TC-002"]
    assert outcome == ms.FAILED and "TimeoutError" in note and "\x1b" not in note
    assert note.startswith("App/environment stall")


def test_the_workbook_has_exactly_one_worksheet_with_filters_and_a_frozen_header(tmp_path):
    rows = _rows({"S-1": [_case("Zoom in"), _case("Pan")]})
    out = ms.write_sheet(rows, tmp_path / "master.xlsx")

    wb = openpyxl.load_workbook(out)
    assert len(wb.sheetnames) == 1
    ws = wb.active
    assert ws.freeze_panes == "F2" and ws.auto_filter.ref
    assert [c.value for c in ws[1]][:5] == ["#", "Story", "Story title", "Test Case ID", "Automation ID"]
    assert ws.max_row == 3


def test_the_summary_counts_add_up():
    rows = _rows({"S-1": [_case("a"), _case("b", execution="Manual")]})
    counts = ms.summarize(rows)
    assert counts["total"] == 2 and counts[ms.AUTOMATABLE] == 1 and counts[ms.MANUAL] == 1


# ---- finding the scripts: ask Playwright, do not grep the source ------------------------------

FAKE_PLAYWRIGHT = '''
import json
print(json.dumps({"suites": [{"file": "e2e/a.spec.ts", "specs": [
  {"title": "PROJ-1-AUTO-TC-094 Grayscale", "file": "e2e/a.spec.ts"},
  {"title": "PROJ-1-AUTO-TC-095 Negative", "file": "e2e/a.spec.ts"}]}]}))
'''


def test_a_test_whose_id_is_built_in_a_loop_still_counts_as_a_script(tmp_path, monkeypatch):
    # The failure this prevents: `test(`PROJ-1-AUTO-${id} ...`)` inside a for-loop has no literal id
    # in the source, so a text scan reported TC-094/095/097/098 as "script not written yet" while
    # Playwright was running them.
    import sys

    cli = tmp_path / "node_modules" / "@playwright" / "test" / "cli.js"
    cli.parent.mkdir(parents=True)
    cli.write_text(FAKE_PLAYWRIGHT, encoding="utf-8")
    (tmp_path / "tests" / "e2e").mkdir(parents=True)
    (tmp_path / "tests" / "e2e" / "a.spec.ts").write_text(
        "for (const id of ['TC-094','TC-095']) { test(`PROJ-1-AUTO-${id} x`, async () => {}); }",
        encoding="utf-8")
    monkeypatch.setattr(ms.shutil if hasattr(ms, "shutil") else __import__("shutil"), "which",
                        lambda name: sys.executable)
    monkeypatch.chdir(tmp_path.parent)          # a RELATIVE workspace path, as the CLI passes it

    found = ms.spec_index(tmp_path / "tests", tmp_path)

    assert set(found) == {"PROJ-1-AUTO-TC-094", "PROJ-1-AUTO-TC-095"}
    assert found["PROJ-1-AUTO-TC-094"] == "tests/e2e/a.spec.ts"


def test_without_node_the_text_scan_still_finds_literal_ids(tmp_path, monkeypatch):
    (tmp_path / "tests").mkdir()
    (tmp_path / "tests" / "a.spec.ts").write_text("test('PROJ-1-AUTO-TC-007 x', async () => {});",
                                                  encoding="utf-8")
    monkeypatch.setattr(__import__("shutil"), "which", lambda name: None)

    assert set(ms.spec_index(tmp_path / "tests", tmp_path)) == {"PROJ-1-AUTO-TC-007"}
