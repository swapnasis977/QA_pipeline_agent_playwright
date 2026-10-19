"""One sheet that answers "what test cases do we have, and which are automated?".

The pipeline produces its cases in pieces, each with its own numbering:

  * per-story sheets   one per JIRA story, ids like `PROJ-68308-TC-001`   (the manual cases)
  * consolidated sheet representative cases picked from those, ids like `PROJ-68294-AUTO-TC-014`
  * the run sheet      what the Coding Agent turned into scripts (may add cases of its own)
  * the spec files     the scripts; a test title starts with its AUTO id
  * run_results.json   how each script did the last time it ran

Nothing joined them, so "how many cases do we have and how many are automated" had no single
answer. This module joins them: one row per case, with an automation status and the last
result. Pure functions over files, so it is testable and can be re-run whenever the suite
changes.
"""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

AUTOMATED = "Automated - script exists"
SELECTED = "Selected for automation - script not written yet"
AUTOMATABLE = "Automatable - no script yet"
MANUAL = "Manual only"

PASSED, FAILED, NOT_RUN = "Passed", "Failed", "Not run"

_AUTO_ID = re.compile(r"PROJ-\d+-AUTO-TC-\d+")
# A hand-written script is titled with the SOURCE case id (PROJ-68479-TC-003) rather than an AUTO id.
_TITLE_ID = re.compile(r"PROJ-\d+-(?:AUTO-)?TC-\d+")
_STORY_IN_FILENAME = re.compile(r"_(PROJ-\d+)_TestCases\.xlsx$")
_ANSI = re.compile(r"\x1b\[[0-9;]*m")

COLUMNS = [
    ("#", 5), ("Story", 11), ("Story title", 34), ("Test Case ID", 22), ("Automation ID", 22),
    ("Name", 46), ("Description", 40), ("Component", 18), ("Sub-Component", 22),
    ("Prerequisite", 32), ("Test Step", 52), ("Test Data", 22), ("Expected Result", 52),
    ("Testing Type", 18), ("Execution Type (as written)", 16), ("Automation Status", 30),
    ("Automated Test File", 44), ("Last Run Result", 13), ("Last Run Note", 48),
    ("Scope Decision", 28), ("Covered By", 24), ("Decision Note", 44),
    ("Product", 16), ("Result", 13), ("Status", 12),
]


def read_rows(path: str | Path) -> list[dict[str, Any]]:
    """The non-blank rows of a test case sheet, as {header: value}."""
    sheet = openpyxl.load_workbook(path, read_only=True, data_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(h or "").strip() for h in rows[0]]
    return [dict(zip(header, r, strict=False)) for r in rows[1:]
            if any(c not in (None, "") for c in r)]


def story_of(path: str | Path) -> str:
    match = _STORY_IN_FILENAME.search(Path(path).name)
    return match.group(1) if match else ""


def spec_index(spec_dir: str | Path, workspace: str | Path) -> dict[str, str]:
    """{AUTO id: workspace-relative spec path} for every test that EXISTS.

    Asks Playwright for the tests it would run (`--list`), because that is the only honest
    answer. Scanning the source text for literal ids missed every test whose id is built at run
    time (`test(`PROJ-68294-AUTO-${id} ...`)` inside a loop), so four cases that had scripts were
    reported as "script not written yet". The text scan is kept as a fallback for a machine
    without Node, and its findings are merged in (a literal id is always a real test).
    """
    found = _listed_by_playwright(workspace)
    root = Path(workspace)
    for spec in sorted(Path(spec_dir).rglob("*.spec.ts")):
        for auto_id in _AUTO_ID.findall(spec.read_text(encoding="utf-8")):
            found.setdefault(auto_id, spec.relative_to(root).as_posix())
    return found


def _listed_by_playwright(workspace: str | Path, timeout: int = 180) -> dict[str, str]:
    """{AUTO id: spec path} from `playwright test --list --reporter=json`, or {} if unavailable."""
    import shutil
    import subprocess

    # Absolute paths: the subprocess runs with cwd=root, so a RELATIVE path to the CLI would stop
    # resolving the moment the directory changes (and the failure was swallowed below).
    root = Path(workspace).resolve()
    cli = root / "node_modules" / "@playwright" / "test" / "cli.js"
    node = shutil.which("node")
    if not (node and cli.is_file()):
        return {}
    try:
        done = subprocess.run([node, str(cli), "test", "--list", "--reporter=json"], cwd=root,
                              capture_output=True, text=True, encoding="utf-8", errors="replace",
                              timeout=timeout)
        report = json.loads(done.stdout)
    except (OSError, subprocess.SubprocessError, ValueError):
        return {}
    found: dict[str, str] = {}

    def visit(suites: list[dict[str, Any]] | None) -> None:
        for suite in suites or []:
            for spec in suite.get("specs") or []:
                match = _TITLE_ID.search(str(spec.get("title") or ""))
                if match:
                    file = str(spec.get("file") or suite.get("file") or "").replace("\\", "/")
                    found.setdefault(match.group(0), f"tests/{file}" if file else "")
            visit(suite.get("suites"))

    visit(report.get("suites"))
    return found


def last_results(run_results: dict[str, Any], env_failures: dict[str, str] | None = None
                 ) -> dict[str, tuple[str, str]]:
    """{AUTO id: (Passed|Failed|Not run, short note)} from a Playwright JSON report.

    A case with several tests takes the WORST outcome, so a half-passing case is not shown
    as passing. A case the report lists without a result (the run was cut off) is "Not run".
    """
    env_failures = env_failures or {}
    rank = {PASSED: 0, NOT_RUN: 1, FAILED: 2}
    out: dict[str, tuple[str, str]] = {}

    def visit(suites: list[dict[str, Any]] | None) -> None:
        for suite in suites or []:
            for spec in suite.get("specs") or []:
                match = _TITLE_ID.search(str(spec.get("title") or ""))
                if not match:
                    continue
                case_id = match.group(0)
                for test in spec.get("tests") or []:
                    results = test.get("results") or []
                    last = results[-1] if results else {}
                    status = last.get("status")
                    if status == "passed":
                        outcome, note = PASSED, ""
                    elif status in ("failed", "timedOut", "interrupted"):
                        message = _ANSI.sub("", str((last.get("error") or {}).get("message") or ""))
                        first = next((ln.strip() for ln in message.splitlines() if ln.strip()), "")
                        outcome, note = FAILED, first[:160]
                        if case_id in env_failures:
                            note = "App/environment stall before the test's own steps. " + note
                    else:
                        outcome, note = NOT_RUN, "Not reached in the last run"
                    if case_id not in out or rank[outcome] > rank[out[case_id][0]]:
                        out[case_id] = (outcome, note)
            visit(suite.get("suites"))

    visit(run_results.get("suites"))
    return out


def automation_status(execution_type: str, auto_id: str, in_suite: bool) -> str:
    """The one-word answer to "is this case automated?"."""
    if auto_id and in_suite:
        return AUTOMATED
    # A case written as Manual stays Manual even if the consolidation happened to pick it as
    # the representative of its group: nobody should be told a manual case is waiting for a script.
    if str(execution_type).strip().lower() == "manual":
        return MANUAL
    if auto_id:
        return SELECTED
    return AUTOMATABLE


def build_rows(
    per_story: dict[str, list[dict[str, Any]]],
    story_titles: dict[str, str],
    consolidated: list[dict[str, Any]],
    run_sheet: list[dict[str, Any]],
    scripts: dict[str, str],
    results: dict[str, tuple[str, str]],
) -> list[dict[str, Any]]:
    """One row per case: every per-story case, plus run-sheet cases the stories never had."""
    # (story, name) -> AUTO id, for the cases the consolidation picked.
    picked: dict[tuple[str, str], str] = {}
    for row in consolidated:
        story = str(row.get("JIRA Ref (Epic/Story)") or "").split("/")[-1].strip()
        picked[(story, str(row.get("Name") or ""))] = str(row.get("Test Case ID") or "")
    picked_names = {name for (_s, name) in picked}

    rows: list[dict[str, Any]] = []

    def add(story: str, source: dict[str, Any], test_id: str, auto_id: str) -> None:
        # A script is found by the case's AUTO id, or by its own id for a hand-written one.
        script_id = auto_id if auto_id in scripts else (test_id if test_id in scripts else "")
        # A last-run result only means something for a case that HAS a script.
        if script_id:
            outcome, note = results.get(script_id, (NOT_RUN, "Not in the last run's report"))
        else:
            outcome, note = "", ""
        rows.append({
            "Story": story, "Story title": story_titles.get(story, ""),
            "Test Case ID": test_id, "Automation ID": auto_id,
            "Name": source.get("Name"), "Description": source.get("Description"),
            "Component": source.get("Component/s"), "Sub-Component": source.get("Sub-Component"),
            "Prerequisite": source.get("Prerequisite"), "Test Step": source.get("Test Step"),
            "Test Data": source.get("Test Data"), "Expected Result": source.get("Expected Result"),
            "Testing Type": source.get("Testing Type"),
            "Execution Type (as written)": source.get("Execution Type"),
            "Automation Status": automation_status(
                str(source.get("Execution Type") or ""), script_id or auto_id, bool(script_id)),
            "Automated Test File": scripts.get(script_id, ""),
            "Last Run Result": outcome,
            "Last Run Note": note,
            "Product": source.get("Product"), "Result": source.get("Result"),
            "Status": source.get("Status"),
            # The untouched source row, for consumers that need every original column
            # (the final automation sheet). Never written to the master workbook.
            "_source": source,
        })

    for story in sorted(per_story):
        for source in per_story[story]:
            auto_id = picked.get((story, str(source.get("Name") or "")), "")
            add(story, source, str(source.get("Test Case ID") or ""), auto_id)

    # Cases that exist only because automation added them (not in any story sheet).
    story_names = {str(r.get("Name") or "") for rows_ in per_story.values() for r in rows_}
    for source in run_sheet:
        name = str(source.get("Name") or "")
        if name in story_names or name in picked_names:
            continue
        ref = str(source.get("JIRA Ref (Epic/Story)") or "").split("/")[-1].strip()
        auto_id = str(source.get("Test Case ID") or "")
        add(ref, source, auto_id, auto_id)
        rows[-1]["Story"] = ref or "(added for automation)"
    return rows


_FILLS = {
    AUTOMATED: "C6EFCE", SELECTED: "FFEB9C", AUTOMATABLE: "FCE4D6", MANUAL: "DDEBF7",
    PASSED: "C6EFCE", FAILED: "FFC7CE", NOT_RUN: "EDEDED",
    "Already automated": "C6EFCE", "Covered by existing script": "E2EFDA",
    "New script needed": "FCE4D6", "Blocked - needs test data or a tool": "FFEB9C",
    "Not on this screen": "EDEDED",
}


def write_sheet(rows: list[dict[str, Any]], out_path: str | Path) -> Path:
    """Exactly ONE worksheet, filterable, with the status columns colour-coded."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Cases"
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for col, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    wrap = Alignment(wrap_text=True, vertical="top")
    for index, row in enumerate(rows, start=1):
        values = [index] + [row.get(name) for name, _w in COLUMNS[1:]]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=index + 1, column=col, value=value)
            cell.alignment = wrap
            colour = _FILLS.get(str(value)) if COLUMNS[col - 1][0] in (
                "Automation Status", "Last Run Result", "Scope Decision") else None
            if colour:
                cell.fill = PatternFill("solid", fgColor=colour)
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(rows) + 1}"
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def summarize(rows: list[dict[str, Any]]) -> dict[str, int]:
    counts: dict[str, int] = {"total": len(rows)}
    for row in rows:
        for key, prefix in (("Automation Status", ""), ("Last Run Result", ""),
                            ("Scope Decision", "Decision: ")):
            label = row.get(key) or ""
            if label:                # a prefix keeps one column's "Manual only" apart from another's
                counts[prefix + label] = counts.get(prefix + label, 0) + 1
    return counts


def load_json(path: str | Path) -> Any:
    try:
        return json.loads(Path(path).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def collect_rows(runs_root: str | Path, workspace: str | Path, feature: str, run: str
                 ) -> list[dict[str, Any]] | None:
    """Every case of a feature, joined with its scripts and last results (see module docstring).

    `runs_root` is the profile's runs folder (runs/<app>). Returns None when the feature has no
    per-story sheets to build from.
    """
    import glob

    feature_dir = Path(runs_root) / f"{feature}_FEATURE"
    run_dir = Path(runs_root) / run
    per_story: dict[str, list[dict[str, Any]]] = {}
    titles: dict[str, str] = {}
    for path in sorted(glob.glob(str(feature_dir / "_internal" / "per_story" / "1_testcases"
                                     / "*_TestCases.xlsx"))):
        story = story_of(path)
        per_story[story] = read_rows(path)
        context = load_json(feature_dir / "_internal" / f"story_context_{story}.json")
        titles[story] = context.get("title", "") if isinstance(context, dict) else ""
    if not per_story:
        return None

    consolidated_path = next(iter(
        glob.glob(str(feature_dir / "1_testcases" / "*Consolidated.xlsx"))
        + glob.glob(str(feature_dir / "_internal" / "*Consolidated.xlsx"))), None)
    run_sheet_path = next(iter(glob.glob(str(run_dir / "1_testcases" / f"{run}_TestCases.xlsx"))), None)
    # The newest report wins: the pipeline run's own, or the latest `execute` of the whole suite.
    candidates = [run_dir / "3_results" / "run_results.json",
                  Path(runs_root) / "SUITE_latest" / "3_results" / "run_results.json"]
    report = load_json(max((c for c in candidates if c.is_file()), key=lambda c: c.stat().st_mtime,
                           default=candidates[0]))
    env = load_json(run_dir / "_internal" / "environment_failures.json")
    results = last_results(report if isinstance(report, dict) else {}, env if isinstance(env, dict) else {})
    return build_rows(
        per_story, titles,
        read_rows(consolidated_path) if consolidated_path else [],
        read_rows(run_sheet_path) if run_sheet_path else [],
        spec_index(Path(workspace) / "tests", workspace), results)
